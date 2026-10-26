"""Carga alumnos por consola y exporta sus notas a un archivo Excel."""

from __future__ import annotations

import math
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from tempfile import NamedTemporaryFile

try:
    import pandas as pd
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError as error:
    print("Faltan dependencias. Ejecutá: pip install -r requirements.txt")
    raise SystemExit(1) from error


NOTA_MINIMA = 0.0
NOTA_MAXIMA = 10.0
PROMEDIO_APROBACION = 6.0
ARCHIVO_SALIDA = Path(__file__).resolve().parent / "reporte_notas.xlsx"


@dataclass
class Alumno:
    """Representa un alumno y las notas ingresadas para él."""

    nombre: str
    notas: list[float]


def validar_nombre(texto: str) -> str:
    """Limpia el nombre y rechaza vacíos, controles y longitud excesiva."""
    nombre = texto.strip()
    if not nombre or len(nombre) > 100:
        raise ValueError("El nombre debe tener entre 1 y 100 caracteres.")
    if any(ord(caracter) < 32 for caracter in nombre):
        raise ValueError("El nombre no puede contener caracteres de control.")
    return nombre


def interpretar_notas(texto: str) -> list[float]:
    """Convierte notas separadas por espacios o ; y valida su rango."""
    partes = re.split(r"[;\s]+", texto.strip())
    if not texto.strip():
        raise ValueError("Ingresá al menos una nota.")
    try:
        notas = [float(parte.replace(",", ".")) for parte in partes]
    except ValueError as error:
        mensaje = "Usá números separados por espacios o punto y coma."
        raise ValueError(mensaje) from error
    if any(not math.isfinite(nota) for nota in notas):
        raise ValueError("Las notas deben ser números finitos.")
    if any(not NOTA_MINIMA <= nota <= NOTA_MAXIMA for nota in notas):
        raise ValueError("Las notas deben estar entre 0 y 10.")
    return notas


def cargar_alumnos() -> list[Alumno]:
    """Solicita alumnos hasta recibir fin; reintenta datos inválidos."""
    alumnos: list[Alumno] = []
    while True:
        texto = input("Nombre del alumno (fin para terminar): ")
        if texto.strip().casefold() == "fin":
            return alumnos
        try:
            nombre = validar_nombre(texto)
        except ValueError as error:
            print(f"Dato inválido: {error}")
            continue
        while True:
            try:
                notas = interpretar_notas(
                    input("Notas (ej.: 7 8,5 6; separá con espacios o ;): ")
                )
            except ValueError as error:
                print(f"Dato inválido: {error}")
                continue
            alumnos.append(Alumno(nombre, notas))
            break


def crear_tabla(alumnos: list[Alumno]) -> pd.DataFrame:
    """Construye la tabla con notas, promedio exacto y condición."""
    if not alumnos:
        raise ValueError("No se cargaron alumnos; no se generó un archivo.")
    filas: list[dict[str, str | float]] = []
    for alumno in alumnos:
        nombre = validar_nombre(alumno.nombre)
        if not alumno.notas:
            raise ValueError(f"El alumno {nombre} no tiene notas.")
        if any(
            not math.isfinite(nota)
            or not NOTA_MINIMA <= nota <= NOTA_MAXIMA
            for nota in alumno.notas
        ):
            raise ValueError(f"El alumno {nombre} tiene notas inválidas.")
        promedio = sum(alumno.notas) / len(alumno.notas)
        fila: dict[str, str | float] = {"Alumno": nombre}
        fila.update(
            {
                f"Nota {indice}": nota
                for indice, nota in enumerate(alumno.notas, 1)
            }
        )
        fila["Promedio"] = promedio
        fila["Condición"] = (
            "Aprobado" if promedio >= PROMEDIO_APROBACION else "Desaprobado"
        )
        filas.append(fila)
    tabla = pd.DataFrame(filas)
    cantidad = max(len(alumno.notas) for alumno in alumnos)
    columnas = ["Alumno"] + [f"Nota {i}" for i in range(1, cantidad + 1)]
    return tabla.reindex(columns=columnas + ["Promedio", "Condición"])


def exportar_excel(tabla: pd.DataFrame, destino: Path) -> None:
    """Guarda una tabla Excel con formato y reemplazo atómico."""
    if tabla.empty:
        raise ValueError("No se puede exportar una tabla vacía.")
    temporal: Path | None = None
    try:
        with NamedTemporaryFile(
            suffix=".xlsx", dir=destino.parent, delete=False
        ) as archivo:
            temporal = Path(archivo.name)
        with pd.ExcelWriter(temporal, engine="openpyxl") as escritor:
            tabla.to_excel(escritor, sheet_name="Notas", index=False)
            hoja = escritor.sheets["Notas"]
            hoja.freeze_panes = "B2"
            for celda in hoja[1]:
                celda.font = Font(bold=True, color="FFFFFF")
                celda.fill = PatternFill("solid", fgColor="24476B")
            for fila in hoja.iter_rows(min_row=2):
                # Los nombres ingresados se guardan como texto, incluso con =.
                fila[0].data_type = "s"
                for celda in fila[1:-1]:
                    celda.number_format = "0.00"
            for columna in hoja.columns:
                ancho = max(len(str(celda.value or "")) for celda in columna)
                hoja.column_dimensions[columna[0].column_letter].width = min(
                    max(ancho + 3, 12), 55
                )
            tabla_excel = Table(displayName="TablaNotas", ref=hoja.dimensions)
            tabla_excel.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=True
            )
            hoja.add_table(tabla_excel)
        temporal.replace(destino)
    finally:
        if temporal is not None:
            temporal.unlink(missing_ok=True)


def mostrar_resumen(tabla: pd.DataFrame, destino: Path) -> None:
    """Muestra el detalle y los totales por condición."""
    aprobados = int((tabla["Condición"] == "Aprobado").sum())
    print("\nReporte generado:")
    print(tabla.to_string(
        index=False, na_rep="—", float_format="{:.2f}".format
    ))
    print(f"\nTotal de alumnos: {len(tabla)}")
    print(f"Aprobados: {aprobados}")
    print(f"Desaprobados: {len(tabla) - aprobados}")
    print(f"Archivo: {destino}")


def main() -> int:
    """Coordina la carga y exportación e informa los errores."""
    print("Reporte de notas | Escala: 0 a 10 | Aprobación: promedio >= 6")
    try:
        alumnos = cargar_alumnos()
        tabla = crear_tabla(alumnos)
        if ARCHIVO_SALIDA.exists():
            respuesta = input("El reporte ya existe. ¿Reemplazarlo? [s/N]: ")
            if respuesta.strip().casefold() != "s":
                print("Guardado cancelado. El archivo anterior se conserva.")
                return 0
        exportar_excel(tabla, ARCHIVO_SALIDA)
        mostrar_resumen(tabla, ARCHIVO_SALIDA)
    except (KeyboardInterrupt, EOFError):
        print("\nCarga interrumpida. No se generó un nuevo reporte.")
        return 1
    except ValueError as error:
        print(f"Aviso: {error}")
        return 1
    except PermissionError:
        print("No se pudo guardar. Cerrá el Excel y revisá los permisos.")
        return 1
    except OSError as error:
        print(f"Error de archivo (ruta, permisos o espacio): {error}")
        return 1
    except Exception as error:
        # Última barrera para informar fallos inesperados en las librerías.
        print(f"Error inesperado ({type(error).__name__}): {error}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
